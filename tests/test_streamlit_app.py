import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

import streamlit_app


class StreamlitAppTests(unittest.TestCase):
    def test_to_number_handles_indonesian_and_plain_amounts(self):
        self.assertEqual(streamlit_app.to_number("Rp 1.234.567"), 1234567)
        self.assertEqual(streamlit_app.to_number("1,234,567"), 1234567)
        self.assertEqual(streamlit_app.to_number("1234,5"), 1234.5)
        self.assertIsNone(streamlit_app.to_number("—"))

    def test_compact_rupiah_uses_short_indonesian_units(self):
        self.assertEqual(streamlit_app.format_compact_rupiah(17_100_000_000), "17,1 M")
        self.assertEqual(streamlit_app.format_compact_rupiah(1_260_000), "1,3 jt")

    def test_account_can_log_in_after_a_new_database_connection(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = Path(directory) / ".data" / "accounts.sqlite3"
            with patch.object(streamlit_app, "ACCOUNT_DB_PATH", database_path):
                salt, digest = streamlit_app.hash_password("password123")
                self.assertTrue(streamlit_app.create_account("staff@example.com", salt, digest))

                account = streamlit_app.get_account("staff@example.com")
                self.assertIsNotNone(account)
                self.assertTrue(streamlit_app.password_matches("password123", account["salt"], account["digest"]))
                self.assertFalse(streamlit_app.password_matches("wrong-password", account["salt"], account["digest"]))
                self.assertFalse(streamlit_app.create_account("staff@example.com", salt, digest))

    def test_revision_compares_remaining_to_increase(self):
        rows = [
            ["", "Semula", None, None, "Menjadi", None, None, None, None],
            [None, "Kode Anggaran", "Uraian", "Jumlah", "Kode Anggaran", "Uraian", "Jumlah", "Selisih", "Sisa Anggaran"],
            [None, None, "Perjalanan", 100, None, "Perjalanan", 120, 20, 15],
            [None, None, "Peralatan", 100, None, "Peralatan", 110, 10, 25],
        ]
        revision = streamlit_app.find_revision(rows)
        self.assertEqual(revision["increase_total"], 30)
        self.assertEqual(revision["insufficient_count"], 1)
        self.assertEqual(revision["sufficient_count"], 1)
        self.assertEqual(revision["shortfall"], 5)

    def test_dashboard_analysis_prepares_rows_and_categories(self):
        rows = [
            ["Kode Anggaran", "Uraian", "Pagu", "Sisa Anggaran"],
            ["051", "Monitoring", 100000, 20000],
            ["052", "Pengawasan", 80000, 10000],
            ["", "Baris tidak bernilai", None, None],
        ]
        analysis = streamlit_app.analyze_rows(rows, "RAB")
        self.assertTrue(analysis["header_found"])
        self.assertEqual(len(analysis["rows"]), 2)
        self.assertEqual(len(analysis["categories"]), 2)
        self.assertEqual(analysis["category_total"], 180000)

    def test_sp2d_pdf_text_parser_reads_category_and_wrapped_detail_rows(self):
        pages = ["""
LAPORAN REALISASI SP2D TA 2026
                       Uraian                    Pagu Revisi Lock Pagu
     100 Kegiatan pertama                     1,000,000 0 500,000 0 500,000 50.00 % 500,000
       101 Kegiatan kedua                      2,000,000 0 1,000,000 0 1,000,000 50.00 % 1,000,000
              000001. Rincian kegiatan yang
                                          100,000 0 50,000 0 50,000 50.00 % 50,000
                 keterangannya terbungkus
"""]

        rows = streamlit_app.parse_sp2d_text(pages)

        self.assertEqual(rows[0], ["Kode Anggaran", "Uraian", "Pagu", "Sisa Anggaran"])
        self.assertEqual(rows[1], ["100", "Kegiatan pertama", 1000000, 500000])
        self.assertEqual(rows[2], ["101", "Kegiatan kedua", 2000000, 1000000])
        self.assertEqual(rows[3], ["000001", "Rincian kegiatan yang keterangannya terbungkus", 100000, 50000])
        self.assertIsNone(streamlit_app.parse_sp2d_text(["PDF laporan lain"]))

    def test_tor_download_is_a_valid_word_document(self):
        analysis = {
            "sheet": "RAB Rev 9",
            "rows": [{
                "Uraian": "Monitoring",
                "Pagu": 100000,
                "Sisa anggaran": 20000,
                "Baris": 2,
                "Kode": "051",
            }],
        }
        sections = streamlit_app.build_tor_sections(
            "Monitoring", "Surabaya", "2026", "", analysis,
        )
        content = streamlit_app.make_tor_docx("Monitoring", "Surabaya", "2026", analysis["sheet"], sections)
        document = Document(io.BytesIO(content))
        self.assertEqual(document.paragraphs[0].text, "KERANGKA ACUAN KERJA")
        self.assertIn("VII. Pembiayaan", [paragraph.text for paragraph in document.paragraphs])
        self.assertTrue(any(paragraph.text.startswith("Sumber data: RAB Rev 9") for paragraph in document.paragraphs))
        self.assertIn("TOR ini menjadi acuan pelaksanaan Monitoring pada tahun 2026, berdasarkan data pada sheet RAB Rev 9.", sections[0][1])
        self.assertLess(max(len(body) for _, body in sections), 180)


if __name__ == "__main__":
    unittest.main()
